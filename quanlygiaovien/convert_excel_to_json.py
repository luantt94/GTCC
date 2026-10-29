"""Đọc các file "BIÊN CHẾ XE KHÓA xxxx.xlsx" (mỗi khóa 1 file, xuất từ hệ
thống quản lý) sang data.json dùng chung cho toàn bộ app.

Mỗi file có 3 sheet:
- "DS": chỉ có số thứ tự, không dùng.
- "Sổ phân công GV": có cột "Thời gian" (Từ ngày... đến ngày...) theo khóa -
  dùng để lấy tu_ngay/den_ngay (sheet chính không còn ghi ngày bắt đầu/kết
  thúc trong tiêu đề như trước).
- Sheet trùng tên mã lớp (vd "0626"): có dòng tiêu đề "BIÊN CHẾ XE KHÓA
  <ma_khoa> (<lop>)" rồi tới bảng học viên, cột E "xe/Gv/Sđt" gộp 3 dòng
  biển số / tên giáo viên / số điện thoại cho nhóm học viên đi theo xe đó.

Không còn sheet lịch học thực hành theo buổi (lịch học giờ nằm ở file PDF
riêng, xem attach_files.py) nên trường "lich_hoc" luôn để trống.
"""
import json
import re
import sys
from datetime import datetime
from pathlib import Path

import openpyxl

TITLE_RE = re.compile(r"BI[ÊE]N CH[ẾE] XE KH[ÓO]A\s+(\S+)\s*\(([^)]+)\)", re.UNICODE | re.IGNORECASE)
PERIOD_RE = re.compile(r"T[ừu]\s*([0-9/]+).*?đ[ếe]n\s*([0-9/]+)", re.UNICODE | re.IGNORECASE | re.DOTALL)


def clean_text(value):
    if value is None:
        return ""
    text = str(value).strip()
    # loại artifact kiểu "...TÊNB18:B20" lẫn vào nội dung do lỗi nhập liệu nguồn
    text = re.sub(r"[A-F]{1,2}\d+:[A-F]{1,2}\d+$", "", text).strip()
    return text


def _find_course_sheet(wb):
    """Tìm sheet có dòng tiêu đề 'BIÊN CHẾ XE KHÓA ...'. Trả về (ws, title_text)."""
    for ws in wb.worksheets:
        for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, 10), max_col=1):
            v = row[0].value
            if v and str(v).strip().upper().startswith("BIÊN CHẾ XE KHÓA"):
                return ws, str(v).strip()
    return None, None


def _parse_groups(ws):
    """Đọc bảng học viên bắt đầu từ dòng có 'Stt' tới hết sheet.
    Trả về list groups: [{giao_vien, xe, sdt, hoc_vien: [...]}]."""
    groups = []
    current_group = None
    started = False

    for row in ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=6):
        a_val = row[0].value
        a_str = str(a_val).strip() if a_val is not None else ""

        if not started:
            if a_str == "Stt":
                started = True
            continue

        if isinstance(a_val, (int, float)):
            stt = int(a_val)
            ho_ten = clean_text(row[1].value)
            ngay_sinh = clean_text(row[2].value)
            dia_chi = clean_text(row[3].value)
            xe_gv_sdt = row[4].value
            ghi_chu = clean_text(row[5].value)

            if xe_gv_sdt and str(xe_gv_sdt).strip():
                lines = [clean_text(x) for x in str(xe_gv_sdt).split("\n")]
                lines = [x for x in lines if x]
                xe = lines[0] if len(lines) > 0 else ""
                giao_vien = lines[1] if len(lines) > 1 else ""
                sdt = lines[2] if len(lines) > 2 else ""
                current_group = {
                    "giao_vien": giao_vien,
                    "xe": xe,
                    "sdt": sdt,
                    "hoc_vien": [],
                }
                groups.append(current_group)

            if current_group is not None:
                current_group["hoc_vien"].append(
                    {
                        "stt": stt,
                        "ho_ten": ho_ten,
                        "ngay_sinh": ngay_sinh,
                        "dia_chi": dia_chi,
                        "ghi_chu": ghi_chu,
                    }
                )

    return groups


def _find_period(wb, ma_khoa):
    """Lấy (tu_ngay, den_ngay) từ sheet 'Sổ phân công GV' theo mã khóa."""
    if "Sổ phân công GV" not in wb.sheetnames:
        return "", ""
    ws = wb["Sổ phân công GV"]
    header_seen = False
    for row in ws.iter_rows(values_only=True):
        if not header_seen:
            if row and str(row[0] or "").strip() == "Số TT":
                header_seen = True
            continue
        if len(row) < 4:
            continue
        khoa = str(row[3] or "").strip()
        if khoa == ma_khoa:
            thoi_gian = str(row[2] or "")
            m = PERIOD_RE.search(thoi_gian)
            if m:
                return m.group(1).strip(), m.group(2).strip()
    return "", ""


def load_course_from_file(path):
    """Đọc 1 file 'BIÊN CHẾ XE KHÓA xxxx.xlsx', trả về 1 course dict."""
    wb = openpyxl.load_workbook(path, data_only=True)
    ws, title = _find_course_sheet(wb)
    if ws is None:
        raise ValueError(
            f"Không tìm thấy dòng tiêu đề 'BIÊN CHẾ XE KHÓA ...' trong file: {path}"
        )

    m = TITLE_RE.search(title)
    if m:
        ma_khoa, lop = m.group(1).strip(), m.group(2).strip()
    else:
        ma_khoa, lop = title, ""

    groups = _parse_groups(ws)
    so_hoc_vien = sum(len(g["hoc_vien"]) for g in groups)
    tu_ngay, den_ngay = _find_period(wb, ma_khoa)

    return {
        "ma_khoa": ma_khoa,
        "lop": lop,
        "tu_ngay": tu_ngay,
        "den_ngay": den_ngay,
        "hv_per_xe": None,
        "so_hoc_vien": so_hoc_vien,
        "groups": groups,
        "lich_hoc": [],  # lịch học giờ chỉ còn ở file PDF, xem attach_files.py
        "source_file": Path(path).name,  # để báo cảnh báo rõ nguồn khi dữ liệu không khớp
    }


def build_data(excel_paths):
    """Đọc & gộp nhiều file 'BIÊN CHẾ XE KHÓA xxxx.xlsx' (mỗi file 1 khóa).
    Nếu 2 file trùng mã khóa, file đọc sau ghi đè file đọc trước (trả thêm
    danh sách mã bị ghi đè)."""
    courses_by_code = {}
    overwritten = []
    for path in excel_paths:
        c = load_course_from_file(path)
        if c["ma_khoa"] in courses_by_code:
            overwritten.append(c["ma_khoa"])
        courses_by_code[c["ma_khoa"]] = c

    courses = [courses_by_code[k] for k in sorted(courses_by_code)]
    data = {
        "generated_at": datetime.now().isoformat(timespec="seconds"),
        "source_files": [Path(p).name for p in excel_paths],
        "courses": courses,
    }
    return data, overwritten


def main():
    paths = [Path(p) for p in sys.argv[1:]]
    if not paths:
        print("Dùng: python convert_excel_to_json.py <file1.xlsx> [file2.xlsx ...]")
        sys.exit(1)
    output_path = Path("data.json")
    data, overwritten = build_data(paths)
    output_path.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"OK: exported {len(data['courses'])} khoa -> {output_path}")
    if overwritten:
        print(f"CANH BAO: trung ma khoa: {overwritten}")


if __name__ == "__main__":
    main()
